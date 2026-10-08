"""Import EVN reservoir snapshots from the user's audit workbook.

Usage: python3 scripts/build_hydro_reservoirs.py INPUT.xlsx
The workbook is an archive, not a live EVN connection. No missing observations are filled.
"""
import calendar
import hashlib
import json
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parents[1] / 'data' / 'hydro-reservoirs.js'
FIELDS = {'level': 9, 'inflow': 12, 'spill': 14, 'generation': 15}

def main(path):
    source = Path(path)
    wb = load_workbook(source, read_only=True, data_only=True)
    rows = defaultdict(lambda: defaultdict(list))
    region = {}
    exact = defaultdict(int)
    observed = defaultdict(set)
    latest = None
    for r in wb['Du_lieu_ngay'].iter_rows(min_row=2, values_only=True):
        requested, name = r[0], r[3]
        if not isinstance(requested, date) or not name:
            continue
        if not isinstance(r[7], (int, float)) or abs(r[7]) > 2:
            continue
        name = 'Khe Bố' if str(name).lower() == 'khe bố' else str(name)
        ym = requested.strftime('%Y-%m')
        key = (name, ym)
        region[name] = r[2]
        observed[key].add(requested.date().isoformat() if hasattr(requested, 'date') else requested.isoformat())
        exact[key] += r[8] is True
        latest = max(latest, requested.date() if hasattr(requested, 'date') else requested) if latest else (requested.date() if hasattr(requested, 'date') else requested)
        for field, idx in FIELDS.items():
            value = r[idx]
            if isinstance(value, (int, float)) and not isinstance(value, bool) and value >= 0:
                rows[key][field].append(float(value))
    data = []
    for (name, ym), fields in sorted(rows.items()):
        y, m = map(int, ym.split('-'))
        count = len(observed[name, ym])
        expected = calendar.monthrange(y, m)[1]
        entry = {'reservoir': name, 'region': region[name], 'month': ym,
                 'days': count, 'expected': expected, 'exact00': exact[name, ym]}
        for field in FIELDS:
            values = fields[field]
            entry[field] = round(sum(values) / len(values), 3) if values else None
            entry[field + '_n'] = len(values)
        data.append(entry)
    payload = {'source': 'EVN reservoir workbook supplied by user',
               'source_url': 'https://www.evn.com.vn/vi-VN/thong-tin-ho-thuy-dien/Muc-nuoc-cac-ho-thuy-dien-60-123',
               'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
               'last_observation': latest.isoformat(), 'sampling': 'one requested 00:00 snapshot per day; retain only EVN timestamps within two hours of 00:00',
               'monthly': data}
    OUT.write_text('window.HYDRO_RESERVOIRS=' + json.dumps(payload, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'{len(region)} reservoirs, {len(data)} reservoir-months, latest {latest}, {OUT}')

if __name__ == '__main__':
    main(sys.argv[1])
