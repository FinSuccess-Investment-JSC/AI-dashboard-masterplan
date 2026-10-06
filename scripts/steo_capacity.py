"""EIA STEO: monthly OPEC capacity and outages (3dtab), and annual world supply/demand (3atab)."""
from __future__ import annotations

import calendar
import datetime as dt
import hashlib
import io
import re

import openpyxl
import requests

URL = 'https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx'
MONTHS = {name: month for month, name in enumerate(calendar.month_abbr) if name}


def parse_workbook(raw: bytes):
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if '3dtab' not in book.sheetnames:
        raise ValueError('EIA STEO 3dtab missing')
    sheet = book['3dtab']
    title = str(sheet.cell(2, 2).value or '')
    if 'Short-Term Energy Outlook' not in title:
        raise ValueError('Wrong EIA STEO workbook')
    match = re.search(r'-\s+([A-Za-z]+)\s+(\d{4})$', title)
    if not match:
        raise ValueError('EIA STEO issue month missing')
    issue = f'{int(match[2]):04d}-{MONTHS[match[1][:3]]:02d}'
    series = {}
    for row in sheet.iter_rows():
        code = row[0].value
        if code in ('cops_opec', 'padi_OPEC'):
            if code in series:
                raise ValueError('Duplicate STEO series identity')
            series[code] = row
    if set(series) != {'cops_opec', 'padi_OPEC'}:
        raise ValueError('EIA STEO OPEC series missing')
    year = None
    records = []
    for col in range(3, sheet.max_column + 1):
        heading = sheet.cell(3, col).value
        if isinstance(heading, (int, float)):
            year = int(heading)
        month = MONTHS.get(str(sheet.cell(4, col).value or '')[:3])
        if not year or not month or year < 2024:
            continue
        capacity = series['cops_opec'][col - 1].value
        outages = series['padi_OPEC'][col - 1].value
        if not isinstance(outages, (int, float)):
            continue  # EIA leaves future outages blank; do not plot forecasts as actuals.
        if not isinstance(capacity, (int, float)) or not -0.5 <= capacity <= 15 or not 0 <= outages <= 20:
            raise ValueError('Invalid STEO OPEC estimate')
        last_day = calendar.monthrange(year, month)[1]
        records.append({'date': f'{year:04d}-{month:02d}-{last_day:02d}',
                        'capacity': round(float(capacity), 4), 'outages': round(float(outages), 4)})
    if len(records) < 24 or records[-1]['date'] > dt.date.today().isoformat():
        raise ValueError('Incomplete or future STEO history')
    return issue, records


def load():
    response = requests.get(URL, timeout=40, headers={'User-Agent': 'FinSuccessDashboard/2.0 (public data research)'})
    response.raise_for_status()
    issue, rows = parse_workbook(response.content)
    return {'source_url': URL, 'source_name': 'EIA STEO', 'observation_frequency': 'monthly',
            'publication_frequency': 'monthly', 'issue': issue, 'unit': 'million barrels/day',
            'records': rows, 'raw_sha256': hashlib.sha256(response.content).hexdigest(),
            'method': '3dtab: cops_opec and padi_OPEC; months with both series only'}, response.content

def parse_world_balance(raw: bytes):
    """Annual means of papr_world (world production) and patc_world (world consumption), 3atab.

    A year with any month at or after the issue month contains EIA forecasts and is flagged
    `estimate` (shown as 2026E). Only years with all 12 months are kept."""
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if '3atab' not in book.sheetnames:
        raise ValueError('EIA STEO 3atab missing')
    sheet = book['3atab']
    title = str(sheet.cell(2, 2).value or '')
    match = re.search(r'-\s+([A-Za-z]+)\s+(\d{4})$', title)
    if 'Short-Term Energy Outlook' not in title or not match:
        raise ValueError('Wrong EIA STEO workbook')
    issue = f'{int(match[2]):04d}-{MONTHS[match[1][:3]]:02d}'
    series = {}
    for row in sheet.iter_rows():
        if row[0].value in ('papr_world', 'patc_world'):
            code = row[0].value
            if code in series and [c.value for c in series[code][2:]] != [c.value for c in row[2:]]:
                raise ValueError('Conflicting STEO series identity')  # 3atab repeats the world total row
            series.setdefault(code, row)
    if set(series) != {'papr_world', 'patc_world'}:
        raise ValueError('EIA STEO world supply/demand missing')
    months, year = {}, None
    for col in range(3, sheet.max_column + 1):
        heading = sheet.cell(3, col).value
        if isinstance(heading, (int, float)):
            year = int(heading)
        month = MONTHS.get(str(sheet.cell(4, col).value or '')[:3])
        if not year or not month:
            continue
        supply, demand = series['papr_world'][col - 1].value, series['patc_world'][col - 1].value
        if not all(isinstance(v, (int, float)) and 80 <= v <= 130 for v in (supply, demand)):
            raise ValueError(f'Invalid STEO world balance {year}-{month:02d}')
        months.setdefault(year, []).append((f'{year:04d}-{month:02d}', supply, demand))
    records = []
    for year, rows in sorted(months.items()):
        if len(rows) != 12:
            continue
        supply = sum(r[1] for r in rows) / 12
        demand = sum(r[2] for r in rows) / 12
        records.append({'date': f'{year:04d}-12-31', 'year': year, 'supply': round(supply, 2), 'demand': round(demand, 2),
                        'balance': round(supply - demand, 2), 'estimate': rows[-1][0] >= issue})
    if len(records) < 4:
        raise ValueError('Incomplete STEO world balance')
    return issue, records


def load_world_balance():
    response = requests.get(URL, timeout=40, headers={'User-Agent': 'FinSuccessDashboard/2.0 (public data research)'})
    response.raise_for_status()
    issue, rows = parse_world_balance(response.content)
    return {'source_url': URL, 'source_name': 'EIA STEO', 'observation_frequency': 'annual (monthly mean)',
            'publication_frequency': 'monthly', 'issue': issue, 'unit': 'million barrels/day',
            'records': rows, 'raw_sha256': hashlib.sha256(response.content).hexdigest(),
            'method': '3atab: papr_world (world total production), patc_world (world total consumption); 12-month mean per year; balance = supply - demand; years with forecast months flagged estimate'}, response.content
