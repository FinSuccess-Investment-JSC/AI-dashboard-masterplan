import csv
import datetime as dt
import io
import sys
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from steo_capacity import parse_workbook, parse_world_balance
from cftc_positions import parse_archive


class PeriodicFeedTests(unittest.TestCase):
    def test_steo_pairs_only_months_with_outages(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '3dtab'
        sheet.cell(2, 2, 'U.S. Energy Information Administration  |  Short-Term Energy Outlook  - September 2026')
        sheet.cell(44, 1, 'cops_opec')
        sheet.cell(49, 1, 'padi_OPEC')
        for offset in range(25):
            col = offset + 3
            year, month = 2024 + offset // 12, offset % 12 + 1
            if month == 1:
                sheet.cell(3, col, year)
            sheet.cell(4, col, dt.date(year, month, 1).strftime('%b'))
            sheet.cell(44, col, 2.0)
            sheet.cell(49, col, 1.0 if offset < 24 else '-  ')
        stream = io.BytesIO()
        book.save(stream)
        issue, records = parse_workbook(stream.getvalue())
        self.assertEqual(issue, '2026-09')
        self.assertEqual(len(records), 24)
        self.assertEqual(records[-1]['date'], '2025-12-31')

    def test_steo_world_balance_annual_means_and_estimates(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '3atab'
        sheet.cell(2, 2, 'U.S. Energy Information Administration  |  Short-Term Energy Outlook  - September 2026')
        for row, code in ((6, 'papr_world'), (10, 'papr_world'), (19, 'patc_world')):  # EIA repeats the world total row
            sheet.cell(row, 1, code)
        for offset in range(60):  # 2022-2026
            col = offset + 3
            year, month = 2022 + offset // 12, offset % 12 + 1
            if month == 1:
                sheet.cell(3, col, year)
            sheet.cell(4, col, dt.date(year, month, 1).strftime('%b'))
            for row in (6, 10):
                sheet.cell(row, col, 100.0 + (year - 2022))
            sheet.cell(19, col, 99.0 + (year - 2022) + (0.6 if month == 12 else 0))
        stream = io.BytesIO()
        book.save(stream)
        issue, records = parse_world_balance(stream.getvalue())
        self.assertEqual(issue, '2026-09')
        self.assertEqual([r['year'] for r in records], [2022, 2023, 2024, 2025, 2026])
        self.assertEqual((records[0]['supply'], records[0]['demand'], records[0]['balance']), (100.0, 99.05, 0.95))
        self.assertEqual([r['estimate'] for r in records], [False, False, False, False, True])
        sheet.cell(10, 5, 50.0)  # a conflicting duplicate row is rejected
        stream = io.BytesIO()
        book.save(stream)
        with self.assertRaises(ValueError):
            parse_world_balance(stream.getvalue())

    def test_cftc_contract_and_net_thousands(self):
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=['Market_and_Exchange_Names', 'CFTC_Contract_Market_Code',
                                                    'Report_Date_as_YYYY-MM-DD', 'M_Money_Positions_Long_All',
                                                    'M_Money_Positions_Short_All'])
        writer.writeheader()
        for week in range(20):
            writer.writerow({'Market_and_Exchange_Names': 'WTI-PHYSICAL - NEW YORK MERCANTILE EXCHANGE',
                             'CFTC_Contract_Market_Code': '067651',
                             'Report_Date_as_YYYY-MM-DD': (dt.date(2025, 1, 7) + dt.timedelta(weeks=week)).isoformat(),
                             'M_Money_Positions_Long_All': '223190', 'M_Money_Positions_Short_All': '121362'})
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, 'w') as z:
            z.writestr('f_year.txt', stream.getvalue())
        rows = parse_archive(archive.getvalue(), 2025)
        self.assertEqual(rows[0]['net_thousands'], 101.828)
        self.assertEqual(len(rows), 20)


if __name__ == '__main__':
    unittest.main()
